from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from docx import Document as WordDocument
from pypdf import PdfReader
from pptx import Presentation


@dataclass(frozen=True)
class ExtractedPage:
    text: str
    page: int | None


def extract_pages(path: Path, resource_type: str) -> list[ExtractedPage]:
    kind = resource_type.upper()
    if kind == "PDF":
        return _extract_pdf(path)
    if kind == "DOC":
        return _extract_docx(path)
    if kind == "PPT":
        return _extract_pptx(path)
    if kind in {"MARKDOWN", "TXT"}:
        return [ExtractedPage(_read_text(path), None)]
    raise ValueError(f"unsupported resource type: {resource_type}")


def _extract_pdf(path: Path) -> list[ExtractedPage]:
    reader = PdfReader(path)
    return [
        ExtractedPage(text, page_number)
        for page_number, page in enumerate(reader.pages, start=1)
        if (text := (page.extract_text() or "").strip())
    ]


def _extract_docx(path: Path) -> list[ExtractedPage]:
    document = WordDocument(path)
    text = "\n".join(paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip())
    return [ExtractedPage(text, None)] if text else []


def _extract_pptx(path: Path) -> list[ExtractedPage]:
    presentation = Presentation(path)
    pages: list[ExtractedPage] = []
    for page_number, slide in enumerate(presentation.slides, start=1):
        text = "\n".join(
            shape.text.strip()
            for shape in slide.shapes
            if hasattr(shape, "text") and shape.text.strip()
        )
        if text:
            pages.append(ExtractedPage(text, page_number))
    return pages


def _read_text(path: Path) -> str:
    data = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return data.decode(encoding).strip()
        except UnicodeDecodeError:
            continue
    raise ValueError("text file encoding is not UTF-8 or GB18030")
